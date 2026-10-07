#!/usr/bin/env python3
"""将 data_track（ik_fail）日志汇总为 Excel。

输出内容：
  1) 直线电机槽位对照
  2) 每帧：各直线电机对应关节的绝对角(deg/rad) + 解算行程(mm)
  3) 解算失败明细：原因、行程、对应角

用法:
  python tools/export_data_track_ik_excel.py logs/data_track_YYYYMMDD_HHMMSS.log
"""
from __future__ import annotations

import json
import math
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# 槽位 L00~L13：与 usb/offline.LINEAR_JOINT_SLOTS 一致
LINEAR_MOTORS: List[Dict[str, Any]] = [
    {"slot": 0, "id": "L00", "bus": "CAN1/FDCAN1", "hw_m": "M1",
     "finger": "ring", "joint": "MCP", "cn": "无名指 MCP", "stroke_ref_mm": 16,
     "short": "ring_MCP", "linkage_f": "F3"},
    {"slot": 1, "id": "L01", "bus": "CAN1/FDCAN1", "hw_m": "M2",
     "finger": "ring", "joint": "DIP", "cn": "无名指 DIP", "stroke_ref_mm": 14,
     "short": "ring_DIP", "linkage_f": "F1"},
    {"slot": 2, "id": "L02", "bus": "CAN1/FDCAN1", "hw_m": "M3",
     "finger": "ring", "joint": "PIP", "cn": "无名指 PIP", "stroke_ref_mm": 12,
     "short": "ring_PIP", "linkage_f": "F2"},
    {"slot": 3, "id": "L03", "bus": "CAN1/FDCAN1", "hw_m": "M4",
     "finger": "pinky", "joint": "MCP", "cn": "小指 MCP", "stroke_ref_mm": 16,
     "short": "pinky_MCP", "linkage_f": "F3"},
    {"slot": 4, "id": "L04", "bus": "CAN1/FDCAN1", "hw_m": "M5",
     "finger": "pinky", "joint": "DIP", "cn": "小指 DIP", "stroke_ref_mm": 14,
     "short": "pinky_DIP", "linkage_f": "F1"},
    {"slot": 5, "id": "L05", "bus": "CAN1/FDCAN1", "hw_m": "M6",
     "finger": "pinky", "joint": "PIP", "cn": "小指 PIP", "stroke_ref_mm": 12,
     "short": "pinky_PIP", "linkage_f": "F2"},
    {"slot": 6, "id": "L06", "bus": "CAN2/FDCAN2", "hw_m": "M1",
     "finger": "index", "joint": "MCP", "cn": "食指 MCP", "stroke_ref_mm": 16,
     "short": "index_MCP", "linkage_f": "F3"},
    {"slot": 7, "id": "L07", "bus": "CAN2/FDCAN2", "hw_m": "M2",
     "finger": "index", "joint": "DIP", "cn": "食指 DIP", "stroke_ref_mm": 14,
     "short": "index_DIP", "linkage_f": "F1"},
    {"slot": 8, "id": "L08", "bus": "CAN2/FDCAN2", "hw_m": "M3",
     "finger": "index", "joint": "PIP", "cn": "食指 PIP", "stroke_ref_mm": 12,
     "short": "index_PIP", "linkage_f": "F2"},
    {"slot": 9, "id": "L09", "bus": "CAN2/FDCAN2", "hw_m": "M4",
     "finger": "mid", "joint": "MCP", "cn": "中指 MCP", "stroke_ref_mm": 16,
     "short": "mid_MCP", "linkage_f": "F3"},
    {"slot": 10, "id": "L10", "bus": "CAN2/FDCAN2", "hw_m": "M5",
     "finger": "mid", "joint": "DIP", "cn": "中指 DIP", "stroke_ref_mm": 14,
     "short": "mid_DIP", "linkage_f": "F1"},
    {"slot": 11, "id": "L11", "bus": "CAN2/FDCAN2", "hw_m": "M6",
     "finger": "mid", "joint": "PIP", "cn": "中指 PIP", "stroke_ref_mm": 12,
     "short": "mid_PIP", "linkage_f": "F2"},
    {"slot": 12, "id": "L12", "bus": "CAN3/FDCAN3", "hw_m": "M1",
     "finger": "thumb", "joint": "MCP", "cn": "拇指 MCP", "stroke_ref_mm": 16,
     "short": "thumb_MCP", "linkage_f": "F3(拇指MCP)"},
    {"slot": 13, "id": "L13", "bus": "CAN3/FDCAN3", "hw_m": "M2",
     "finger": "thumb", "joint": "PIP", "cn": "拇指 PIP", "stroke_ref_mm": 12,
     "short": "thumb_PIP", "linkage_f": "F2(拇指PIP)"},
]

# 失败记录中 motor 编号 → 连杆 F 与关节键
_MOTOR_TO_JOINT_KEY = {1: "dip", 2: "pip", 3: "mcp"}  # 四指 F1/F2/F3
_FINGER_CN = {
    "thumb": "拇指", "index": "食指", "mid": "中指", "ring": "无名指", "pinky": "小指",
}

_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(color="FFFFFF", bold=True, name="微软雅黑", size=10)
_FAIL_FILL = PatternFill("solid", fgColor="FFF2CC")
_ALT_FILL = PatternFill("solid", fgColor="F2F2F2")
_THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _load_records(path: Path) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        rows.append(json.loads(line))
    return rows


def _style_header(ws, row: int = 1) -> None:
    for cell in ws[row]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _CENTER
        cell.border = _THIN


def _autosize(ws, min_w: int = 8, max_w: int = 36) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = min_w
        for cell in col:
            if cell.value is None:
                continue
            width = max(width, min(max_w, len(str(cell.value)) + 2))
        ws.column_dimensions[letter].width = width


def _deg(rad: Optional[float]) -> Optional[float]:
    if rad is None:
        return None
    return round(math.degrees(float(rad)), 4)


def _fail_joint_info(fail: Dict[str, Any]) -> Dict[str, Any]:
    """从一条 ik_failure 抽出失败电机对应的关节与角度。"""
    finger = str(fail.get("finger") or "")
    ctx = fail.get("context") or {}
    motor = ctx.get("motor")  # 1/2/3 = F1/F2/F3
    raw_deg = fail.get("raw_deg") or {}
    raw_rad = fail.get("raw_rad") or {}

    if finger == "thumb":
        # 拇指：mcp_motor / pip_motor，context.motor 为 1-based F 编号
        joint_key = None
        if motor == 1 or motor == "1":
            # 默认拇指 MCP 用 F3 系；若仅有 mcp/pip 键则按 phi 对照
            joint_key = "mcp" if "mcp" in raw_deg else "pip"
        elif motor == 2 or motor == "2":
            joint_key = "pip"
        elif motor == 3 or motor == "3":
            joint_key = "mcp"
        else:
            joint_key = "mcp" if "mcp" in raw_deg else "pip"
    else:
        joint_key = _MOTOR_TO_JOINT_KEY.get(int(motor) if motor is not None else -1)

    short = f"{finger}_{joint_key.upper()}" if joint_key else finger
    # 槽位查找
    slot_meta = next((m for m in LINEAR_MOTORS if m["short"] == short), None)
    if slot_meta is None and joint_key:
        # pinky_MCP 等大小写
        slot_meta = next(
            (m for m in LINEAR_MOTORS
             if m["finger"] == finger and m["joint"].lower() == joint_key),
            None,
        )

    angle_deg = raw_deg.get(joint_key) if joint_key else None
    angle_rad = raw_rad.get(joint_key) if joint_key else None
    return {
        "finger": finger,
        "finger_cn": _FINGER_CN.get(finger, finger),
        "joint_key": joint_key,
        "short": short if slot_meta is None else slot_meta["short"],
        "slot_meta": slot_meta,
        "angle_deg": angle_deg,
        "angle_rad": angle_rad,
        "phi_in_deg": ctx.get("phi_in_deg"),
        "reachable": ctx.get("reachable_phi_deg"),
        "motor_f": f"F{motor}" if motor is not None else "",
    }


def _write_ref_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("01_直线电机对照")
    headers = [
        "槽位", "电机ID", "总线", "硬件电机", "手指", "关节", "中文",
        "关节短名", "连杆电机", "90°参考行程(mm)",
        "说明",
    ]
    ws.append(headers)
    for m in LINEAR_MOTORS:
        ws.append([
            m["slot"], m["id"], m["bus"], m["hw_m"], m["finger"], m["joint"],
            m["cn"], m["short"], m["linkage_f"], m["stroke_ref_mm"],
            "行程=连杆逆解(CF mm)；角度=UI关节绝对角(相对零位)",
        ])
    _style_header(ws)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.border = _THIN
            cell.alignment = _CENTER
    note = ws.max_row + 2
    ws.cell(note, 1, "连杆约定：四指 F1=DIP，F2=PIP，F3=MCP；行程列 linear_mm[slot] 与 L00~L13 一一对应。")
    ws.merge_cells(start_row=note, start_column=1, end_row=note, end_column=8)
    _autosize(ws, max_w=42)
    ws.freeze_panes = "A2"


def _write_angles_sheet(wb: Workbook, records: List[Dict[str, Any]]) -> None:
    """宽表：每帧一行；每个直线电机一组 绝对角deg / 绝对角rad / 行程mm。"""
    ws = wb.create_sheet("02_直线电机角度行程")

    # 两行表头：组名 + 子列
    meta_headers = [
        "seq", "时间", "kind", "source", "hand_side", "失败指数",
    ]
    # 第一行组标题
    row1: List[Any] = list(meta_headers)
    row2: List[Any] = [""] * len(meta_headers)
    merges: List[tuple] = []
    col = len(meta_headers) + 1
    for m in LINEAR_MOTORS:
        title = f"{m['id']}|{m['cn']}|{m['short']}|{m['linkage_f']}"
        row1.extend([title, "", ""])
        row2.extend(["绝对角(deg)", "绝对角(rad)", "解算行程(mm)"])
        merges.append((col, col + 2))
        col += 3

    ws.append(row1)
    ws.append(row2)
    for c1, c2 in merges:
        ws.merge_cells(start_row=1, start_column=c1, end_row=1, end_column=c2)

    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _CENTER
        cell.border = _THIN
    for cell in ws[2]:
        cell.fill = PatternFill("solid", fgColor="2E75B6")
        cell.font = _HEADER_FONT
        cell.alignment = _CENTER
        cell.border = _THIN

    fail_shorts = set()
    for rec in records:
        angles = rec.get("angles_by_name") or {}
        linear = list(rec.get("linear_mm") or [])
        fails = rec.get("ik_failures") or []
        fail_shorts.clear()
        for f in fails:
            info = _fail_joint_info(f)
            fail_shorts.add(info["short"])

        row: List[Any] = [
            rec.get("seq"),
            rec.get("t"),
            rec.get("kind"),
            rec.get("source"),
            rec.get("hand_side"),
            len(fails),
        ]
        for m in LINEAR_MOTORS:
            short = m["short"]
            rad = angles.get(short)
            try:
                rad_f = float(rad) if rad is not None else None
            except (TypeError, ValueError):
                rad_f = None
            stroke = linear[m["slot"]] if m["slot"] < len(linear) else None
            row.extend([
                _deg(rad_f) if rad_f is not None else None,
                round(rad_f, 6) if rad_f is not None else None,
                stroke,
            ])
        ws.append(row)
        r = ws.max_row
        for c, cell in enumerate(ws[r], start=1):
            cell.border = _THIN
            cell.alignment = _CENTER if c > 2 else _LEFT
            if r % 2 == 0:
                cell.fill = _ALT_FILL
        # 高亮失败电机三列
        base = len(meta_headers) + 1
        for i, m in enumerate(LINEAR_MOTORS):
            if m["short"] not in fail_shorts:
                continue
            for off in range(3):
                cell = ws.cell(r, base + i * 3 + off)
                cell.fill = _FAIL_FILL

    ws.freeze_panes = "G3"
    # 列宽
    for i, h in enumerate(meta_headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, min(22, len(h) + 4))
    for i in range(len(meta_headers) + 1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(i)].width = 12
    ws.row_dimensions[1].height = 32


def _write_long_sheet(wb: Workbook, records: List[Dict[str, Any]]) -> None:
    """长表：一帧×一电机一行，便于筛选对齐。"""
    ws = wb.create_sheet("03_电机关节长表")
    headers = [
        "seq", "时间", "kind", "source", "hand_side",
        "槽位", "电机ID", "总线", "手指", "关节", "中文", "关节短名", "连杆",
        "绝对角(deg)", "绝对角(rad)", "解算行程(mm)", "本帧是否失败电机",
    ]
    ws.append(headers)
    _style_header(ws)

    for rec in records:
        angles = rec.get("angles_by_name") or {}
        linear = list(rec.get("linear_mm") or [])
        fails = rec.get("ik_failures") or []
        fail_shorts = {_fail_joint_info(f)["short"] for f in fails}
        for m in LINEAR_MOTORS:
            short = m["short"]
            rad = angles.get(short)
            try:
                rad_f = float(rad) if rad is not None else None
            except (TypeError, ValueError):
                rad_f = None
            stroke = linear[m["slot"]] if m["slot"] < len(linear) else None
            is_fail = short in fail_shorts
            row = [
                rec.get("seq"), rec.get("t"), rec.get("kind"),
                rec.get("source"), rec.get("hand_side"),
                m["slot"], m["id"], m["bus"], m["finger"], m["joint"],
                m["cn"], short, m["linkage_f"],
                _deg(rad_f) if rad_f is not None else None,
                round(rad_f, 6) if rad_f is not None else None,
                stroke,
                "是" if is_fail else "否",
            ]
            ws.append(row)
            r = ws.max_row
            for cell in ws[r]:
                cell.border = _THIN
                cell.alignment = _CENTER
                if is_fail:
                    cell.fill = _FAIL_FILL
                elif r % 2 == 0:
                    cell.fill = _ALT_FILL

    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    ws.freeze_panes = "A2"
    _autosize(ws, max_w=28)


def _write_fail_sheet(wb: Workbook, records: List[Dict[str, Any]]) -> None:
    ws = wb.create_sheet("04_解算失败明细")
    headers = [
        "seq", "时间", "kind", "source", "hand_side", "space",
        "失败手指", "失败连杆电机", "失败关节", "电机槽位", "电机ID", "中文关节",
        "解算失误原因",
        "目标phi(deg)", "可达phi下限(deg)", "可达phi上限(deg)",
        "失败关节绝对角(deg)", "失败关节绝对角(rad)",
        "同指MCP(deg)", "同指PIP(deg)", "同指DIP(deg)",
        "解算行程_该电机(mm)",
        "L00行程", "L01", "L02", "L03", "L04", "L05",
        "L06", "L07", "L08", "L09", "L10", "L11", "L12", "L13",
        "是否越界钳位",
    ]
    ws.append(headers)
    _style_header(ws)

    for rec in records:
        angles = rec.get("angles_by_name") or {}
        linear = list(rec.get("linear_mm") or [])
        while len(linear) < 14:
            linear.append(None)
        for fail in rec.get("ik_failures") or []:
            info = _fail_joint_info(fail)
            meta = info["slot_meta"]
            raw_deg = fail.get("raw_deg") or {}
            reachable = info["reachable"] or [None, None]
            lo = reachable[0] if len(reachable) > 0 else None
            hi = reachable[1] if len(reachable) > 1 else None
            slot = meta["slot"] if meta else None
            stroke = linear[slot] if slot is not None and slot < len(linear) else None
            reason = fail.get("reason") or ""
            row = [
                rec.get("seq"),
                rec.get("t"),
                rec.get("kind"),
                rec.get("source"),
                rec.get("hand_side"),
                fail.get("space"),
                info["finger_cn"],
                info["motor_f"],
                (info["joint_key"] or "").upper(),
                slot if slot is not None else "",
                meta["id"] if meta else "",
                meta["cn"] if meta else "",
                reason,
                round(float(info["phi_in_deg"]), 6) if info["phi_in_deg"] is not None else None,
                round(float(lo), 6) if lo is not None else None,
                round(float(hi), 6) if hi is not None else None,
                info["angle_deg"],
                info["angle_rad"],
                raw_deg.get("mcp"),
                raw_deg.get("pip"),
                raw_deg.get("dip"),
                stroke,
                *linear[:14],
                "是" if fail.get("out_of_range") else "否",
            ]
            ws.append(row)
            r = ws.max_row
            for cell in ws[r]:
                cell.border = _THIN
                cell.alignment = _CENTER if cell.column != 13 else _LEFT
                cell.fill = _FAIL_FILL

    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    ws.freeze_panes = "A2"
    _autosize(ws, max_w=48)
    ws.column_dimensions["M"].width = 56  # 原因列


def _write_summary(wb: Workbook, records: List[Dict[str, Any]], src: Path) -> None:
    ws = wb.create_sheet("00_汇总说明")
    fail_n = sum(len(r.get("ik_failures") or []) for r in records)
    by_finger: Dict[str, int] = {}
    by_motor: Dict[str, int] = {}
    for r in records:
        for f in r.get("ik_failures") or []:
            info = _fail_joint_info(f)
            by_finger[info["finger_cn"]] = by_finger.get(info["finger_cn"], 0) + 1
            key = f"{info['motor_f']}/{(info['joint_key'] or '').upper()}"
            by_motor[key] = by_motor.get(key, 0) + 1

    lines = [
        ["字段", "内容"],
        ["源日志", str(src)],
        ["失败帧数", len(records)],
        ["失败条目数(ik_failures)", fail_n],
        ["说明1", "绝对角 = UI 关节角，相对机械零位（非相对邻指）。单位 deg/rad。"],
        ["说明2", "解算行程 = joint_angles→连杆逆解后的 CF 行程(mm)，对应 USB linear_mm[L00~L13]。"],
        ["说明3", "黄色高亮 = 该帧/该电机出现逆解失败。"],
        ["说明4", "四指连杆：F1↔DIP，F2↔PIP，F3↔MCP。"],
        [],
        ["失败手指统计", "次数"],
    ]
    for k, v in sorted(by_finger.items(), key=lambda x: -x[1]):
        lines.append([k, v])
    lines.append([])
    lines.append(["失败连杆电机/关节", "次数"])
    for k, v in sorted(by_motor.items(), key=lambda x: -x[1]):
        lines.append([k, v])
    lines.append([])
    lines.append(["工作表", "内容"])
    lines.append(["01_直线电机对照", "L00~L13 与关节/总线一一对应"])
    lines.append(["02_直线电机角度行程", "宽表：每帧各电机绝对角+行程，失败电机黄底"])
    lines.append(["03_电机关节长表", "长表：一帧一电机一行，可筛选"])
    lines.append(["04_解算失败明细", "原因、目标phi、可达范围、对应角、解算行程"])

    for row in lines:
        ws.append(row)
    _style_header(ws)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=2):
        for cell in row:
            cell.border = _THIN
            cell.alignment = _LEFT
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 88


def export(log_path: Path, out_path: Optional[Path] = None) -> Path:
    records = _load_records(log_path)
    if out_path is None:
        out_path = log_path.with_name(log_path.stem + "_ik_fail.xlsx")

    wb = Workbook()
    # 删默认 sheet
    default = wb.active
    wb.remove(default)

    _write_ref_sheet(wb)
    _write_angles_sheet(wb, records)
    _write_long_sheet(wb, records)
    _write_fail_sheet(wb, records)
    _write_summary(wb, records, log_path)
    # 保证「00_汇总说明」在最前
    wb.move_sheet("00_汇总说明", offset=-len(wb.sheetnames) + 1)

    wb.save(out_path)
    return out_path


def main() -> int:
    if len(sys.argv) < 2:
        print("用法: python tools/export_data_track_ik_excel.py <data_track.log> [out.xlsx]")
        return 2
    log_path = Path(sys.argv[1]).resolve()
    out = Path(sys.argv[2]).resolve() if len(sys.argv) > 2 else None
    if not log_path.is_file():
        print(f"找不到日志: {log_path}")
        return 1
    path = export(log_path, out)
    print(f"已写出: {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
